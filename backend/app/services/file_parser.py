# import os
# import fitz
# import docx
# from fastapi import UploadFile, HTTPException


# UPLOAD_DIR = "uploads"


# def save_upload_file(file: UploadFile) -> str:
#     os.makedirs(UPLOAD_DIR, exist_ok=True)

#     file_path = os.path.join(UPLOAD_DIR, file.filename)

#     with open(file_path, "wb") as buffer:
#         buffer.write(file.file.read())

#     return file_path


# def extract_text_from_pdf(file_path: str) -> str:
#     text = ""

#     try:
#         doc = fitz.open(file_path)

#         for page in doc:
#             text += page.get_text()

#         doc.close()

#     except Exception:
#         raise HTTPException(
#             status_code=400,
#             detail="Unable to parse PDF file"
#         )

#     return text.strip()


# def extract_text_from_docx(file_path: str) -> str:
#     text = ""

#     try:
#         document = docx.Document(file_path)

#         for paragraph in document.paragraphs:
#             text += paragraph.text + "\n"

#     except Exception:
#         raise HTTPException(
#             status_code=400,
#             detail="Unable to parse DOCX file"
#         )

#     return text.strip()


# def extract_resume_text(file_path: str) -> str:
#     lower_path = file_path.lower()

#     if lower_path.endswith(".pdf"):
#         return extract_text_from_pdf(file_path)

#     if lower_path.endswith(".docx"):
#         return extract_text_from_docx(file_path)

#     raise HTTPException(
#         status_code=400,
#         detail="Only PDF and DOCX files are supported"
#     )

import os
import uuid
import fitz
import docx
from fastapi import UploadFile, HTTPException


UPLOAD_DIR = "uploads"


def save_upload_file(file: UploadFile, user_id: str) -> str:
    os.makedirs(UPLOAD_DIR, exist_ok=True)

    original_filename = file.filename.replace(" ", "_")
    unique_filename = f"{user_id}_{uuid.uuid4()}_{original_filename}"

    file_path = os.path.join(UPLOAD_DIR, unique_filename)

    with open(file_path, "wb") as buffer:
        buffer.write(file.file.read())

    return file_path


def extract_text_from_pdf(file_path: str) -> str:
    text = ""

    try:
        doc = fitz.open(file_path)

        for page in doc:
            text += page.get_text()

        doc.close()

    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Unable to parse PDF file"
        )

    return text.strip()


def extract_text_from_docx(file_path: str) -> str:
    text = ""

    try:
        document = docx.Document(file_path)

        for paragraph in document.paragraphs:
            text += paragraph.text + "\n"

    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Unable to parse DOCX file"
        )

    return text.strip()


def extract_resume_text(file_path: str) -> str:
    lower_path = file_path.lower()

    if lower_path.endswith(".pdf"):
        return extract_text_from_pdf(file_path)

    if lower_path.endswith(".docx"):
        return extract_text_from_docx(file_path)

    raise HTTPException(
        status_code=400,
        detail="Only PDF and DOCX files are supported"
    )